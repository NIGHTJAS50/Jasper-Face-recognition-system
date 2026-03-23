from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

doc = Document()

# ── Page margins ──────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin   = Cm(3.17)
    section.right_margin  = Cm(3.17)

# ── Styles ─────────────────────────────────────────────────────────────────────
normal_style = doc.styles['Normal']
normal_style.font.name = 'Times New Roman'
normal_style.font.size = Pt(12)

for i in range(1, 5):
    h = doc.styles[f'Heading {i}']
    h.font.name = 'Times New Roman'
    h.font.bold = True
    h.font.color.rgb = RGBColor(0, 0, 0)

doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(14)
doc.styles['Heading 3'].font.size = Pt(13)
doc.styles['Heading 4'].font.size = Pt(12)

# ── Screenshot placeholders ────────────────────────────────────────────────────
# Each entry: (trigger_text_substring, figure_number, caption)
SCREENSHOT_TRIGGERS = {
    "Everything runs on the phone. No server, no internet, no cloud.": (
        1, "Home screen — registered users list"
    ),
    "Registering a New Person": (
        2, "Registration screen — camera capturing a face with capture progress"
    ),
    "Matching the Face to a Registered Person": (
        3, "Recognition screen — green bounding box with name and confidence %"
    ),
    "Full validation — welcome screen appears": (
        4, "Welcome animation — full-screen card displayed after 95%+ confidence"
    ),
    "Attendance Tracking": (
        5, "Attendance mode — live roll-call list showing present count"
    ),
    "Analytics": (
        7, "Analytics dashboard — four stat cards showing today's numbers"
    ),
    "App Architecture": (
        8, "Settings screen — threshold slider and configuration options"
    ),
}

# Kiosk placeholder inserted manually after attendance section trigger
KIOSK_TRIGGER = "Attendance Tracking"

def set_cell_bg(cell, hex_color):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement('w:shd')
    shd.set(qn('w:val'),   'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'),  hex_color)
    tcPr.append(shd)

def add_screenshot_placeholder(fig_num, caption):
    """Insert a grey bordered box as a screenshot placeholder."""
    # 1-row, 1-col table acting as the image box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell = tbl.rows[0].cells[0]
    # Set minimum row height to ~5 cm
    tr = tbl.rows[0]._tr
    trPr = OxmlElement('w:trPr')
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), '2835')   # ~5 cm in twentieths of a point
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)
    tr.insert(0, trPr)

    # Set cell width to ~14 cm
    from docx.oxml.ns import qn as _qn
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(_qn('w:w'), '7938')   # ~14 cm
    tcW.set(_qn('w:type'), 'dxa')
    tcPr.append(tcW)

    set_cell_bg(cell, 'D9D9D9')

    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f'[ Screenshot — Figure {fig_num} ]')
    run.font.name = 'Times New Roman'
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = RGBColor(80, 80, 80)

    # Caption below the box
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(f'Figure {fig_num}: {caption}')
    cap_run.font.name = 'Times New Roman'
    cap_run.font.size = Pt(10)
    cap_run.italic = True
    cap.paragraph_format.space_after = Pt(12)

def add_bold_inline(para, text):
    parts = re.split(r'\*\*(.*?)\*\*', text)
    for idx, part in enumerate(parts):
        run = para.add_run(part)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        if idx % 2 == 1:
            run.bold = True

def add_paragraph(text, style='Normal', align=None):
    p = doc.add_paragraph(style=style)
    add_bold_inline(p, text)
    if align:
        p.alignment = align
    p.paragraph_format.space_after  = Pt(6)
    p.paragraph_format.space_before = Pt(0)
    return p

def add_code_block(lines):
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.name = 'Courier New'
        run.font.size = Pt(9)
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  'F2F2F2')
        pPr.append(shd)

def add_table_from_md(rows):
    if len(rows) < 2:
        return
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ''
            p = cell.paragraphs[0]
            run = p.add_run(cell_text.strip())
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            if r_idx == 0:
                run.bold = True
                set_cell_bg(cell, 'D9D9D9')
    doc.add_paragraph()

# ── Read source ────────────────────────────────────────────────────────────────
with open(r'c:\Users\HP\Desktop\Project J\docs\JASPER_ACADEMIC_ARTICLE.md',
          encoding='utf-8') as f:
    lines = f.readlines()

# ── Parse and render ───────────────────────────────────────────────────────────
i = 0
in_code  = False
in_table = False
code_lines  = []
table_rows  = []
kiosk_inserted = False

def flush_code():
    global code_lines
    if code_lines:
        add_code_block(code_lines)
        doc.add_paragraph()
        code_lines = []

def flush_table():
    global table_rows
    clean = [r for r in table_rows if not re.match(r'^\s*[\|\-\s]+$', r[0])]
    if len(clean) >= 2:
        add_table_from_md(clean)
    table_rows = []

while i < len(lines):
    raw  = lines[i].rstrip('\n')
    line = raw.strip()

    # ── code fence ──
    if line.startswith('```'):
        if not in_code:
            in_code    = True
            code_lines = []
        else:
            in_code = False
            flush_code()
        i += 1
        continue

    if in_code:
        code_lines.append(raw)
        i += 1
        continue

    # ── table row ──
    if line.startswith('|'):
        cells = [c for c in line.split('|') if c != '']
        if not all(re.match(r'^[\-\s:]+$', c.strip()) for c in cells):
            table_rows.append([c.strip() for c in cells])
        i += 1
        in_table = True
        continue
    else:
        if in_table:
            flush_table()
            in_table = False

    # ── headings ──
    if line.startswith('# ') and not line.startswith('## '):
        flush_code()
        h = doc.add_heading(line[2:], level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        i += 1
        continue

    if line.startswith('## '):
        flush_code()
        heading_text = line[3:]
        doc.add_heading(heading_text, level=2)
        # Insert screenshot after heading if triggered
        for trigger, (fig_num, caption) in SCREENSHOT_TRIGGERS.items():
            if trigger in heading_text:
                add_screenshot_placeholder(fig_num, caption)
                break
        i += 1
        continue

    if line.startswith('### '):
        flush_code()
        heading_text = line[4:]
        doc.add_heading(heading_text, level=3)
        for trigger, (fig_num, caption) in SCREENSHOT_TRIGGERS.items():
            if trigger in heading_text:
                add_screenshot_placeholder(fig_num, caption)
                # Insert kiosk placeholder right after attendance placeholder
                if trigger == KIOSK_TRIGGER and not kiosk_inserted:
                    add_screenshot_placeholder(6, "Kiosk mode — green AUTHORIZED screen after a recognized face")
                    kiosk_inserted = True
                break
        i += 1
        continue

    if line.startswith('#### '):
        flush_code()
        doc.add_heading(line[5:], level=4)
        i += 1
        continue

    # ── horizontal rule ──
    if line.startswith('---'):
        doc.add_paragraph('─' * 60)
        i += 1
        continue

    # ── blank line ──
    if line == '':
        i += 1
        continue

    # ── bold-only lines ──
    if line.startswith('**') and line.endswith('**') and line.count('**') == 2:
        p = doc.add_paragraph()
        run = p.add_run(line[2:-2])
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        i += 1
        continue

    # ── regular paragraph — check for screenshot triggers ──
    para_added = False
    for trigger, (fig_num, caption) in SCREENSHOT_TRIGGERS.items():
        if trigger in line:
            add_paragraph(line)
            add_screenshot_placeholder(fig_num, caption)
            if trigger == KIOSK_TRIGGER and not kiosk_inserted:
                add_screenshot_placeholder(6, "Kiosk mode — green AUTHORIZED screen after a recognized face")
                kiosk_inserted = True
            para_added = True
            break
    if not para_added:
        add_paragraph(line)
    i += 1

# flush anything remaining
if in_code:
    flush_code()
if in_table:
    flush_table()

# ── Save ───────────────────────────────────────────────────────────────────────
out_path = r'c:\Users\HP\Desktop\Project J\docs\JASPER_ARTICLE_FINAL.docx'  # overwrite the preferred file
doc.save(out_path)
print(f"Saved: {out_path}")
